import base64
import io
import json
import os
from unittest.mock import patch

import openpyxl

from api import analyze as analyze_mod


class _FakeRequest:
    """Minimal stand-in for the parts of BaseHTTPRequestHandler that do_POST touches."""

    def __init__(self, body: bytes):
        self.rfile = io.BytesIO(body)
        self.headers = {"Content-Length": str(len(body))}
        self.wfile = io.BytesIO()
        self.status = None

    def send_response(self, status):
        self.status = status

    def send_header(self, *_args, **_kwargs):
        pass

    def end_headers(self):
        pass


def _post(payload: dict, env: dict | None = None) -> tuple[int, dict]:
    body = json.dumps(payload).encode()
    req = _FakeRequest(body)
    handler = analyze_mod.handler.__new__(analyze_mod.handler)
    handler.rfile = req.rfile
    handler.headers = req.headers
    handler.wfile = req.wfile
    handler.send_response = req.send_response
    handler.send_header = req.send_header
    handler.end_headers = req.end_headers

    with patch.dict(os.environ, env or {}, clear=env is not None):
        handler.do_POST()

    return req.status, json.loads(req.wfile.getvalue())


def test_missing_api_key_returns_500():
    status, data = _post({"content": "a,b\n1,2\n", "filename": "x.csv"}, env={})
    assert status == 500
    assert "GEMINI_API_KEY" in data["error"]


def test_missing_content_returns_400():
    status, data = _post({"filename": "x.csv"}, env={"GEMINI_API_KEY": "fake-key"})
    assert status == 400
    assert "content" in data["error"]


def test_invalid_config_value_returns_400():
    status, data = _post(
        {"content": "a,b\n1,2\n", "tool": "NotAValidTool"},
        env={"GEMINI_API_KEY": "fake-key"},
    )
    assert status == 400
    assert "preferred_tool" in data["error"]


_MOCK_PLAN = {
    "dataset_classification": "Test",
    "executive_summary": "Column a totals 3 across two rows.",
    "kpis": [{"name": "Total A", "column": "a", "agg": "sum", "format": "number"}],
    "charts": [{"title": "A by B", "type": "bar", "category_column": "b", "value_column": "a", "agg": "sum"}],
    "insights_text": ["a is small but consistent."],
    "risks": [],
    "opportunities": [],
    "recommendations": [],
    "assumptions": [],
    "limitations": [],
}


def test_happy_path_returns_summary_and_workbook():
    with patch.dict(os.environ, {"GEMINI_API_KEY": "fake-key"}):
        with patch.object(analyze_mod.DataAnalystAgent, "plan_dashboard", return_value=_MOCK_PLAN):
            status, data = _post(
                {"content": "a,b\n1,x\n2,y\n", "filename": "sales.csv", "depth": "Quick Analysis"}
            )

    assert status == 200
    assert data["workbook_filename"] == "sales_dashboard.xlsx"
    assert data["summary"]["dataset_classification"] == "Test"
    assert data["summary"]["kpis"] == [{"name": "Total A", "value": "3", "meaning": ""}]

    workbook_bytes = base64.b64decode(data["workbook_base64"])
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
    assert "Overview" in wb.sheetnames
    assert "Data" in wb.sheetnames
