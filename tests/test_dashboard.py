from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from datamind.dashboard import (
    build_workbook,
    compute_kpi_summary,
    compute_kpi_value,
    format_kpi_value,
    load_dataframe,
)

SAMPLE_PATH = Path(__file__).parent.parent / "sample_data" / "sample_sales.csv"


@pytest.fixture
def df():
    return load_dataframe(SAMPLE_PATH)


@pytest.fixture
def plan():
    return {
        "dataset_classification": "Sales",
        "executive_summary": "Revenue grew steadily, led by the North region.",
        "filter_columns": ["region", "category", "does_not_exist", "revenue"],
        "kpis": [
            {"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency", "meaning": "Total sales"},
            {"name": "Total Orders", "agg": "row_count", "format": "number", "meaning": "Number of orders"},
            {"name": "Avg Discount", "column": "discount", "agg": "mean", "format": "percent", "meaning": "Avg discount"},
            {"name": "Bogus", "column": "does_not_exist", "agg": "sum", "format": "number"},
        ],
        "charts": [
            {"title": "Revenue by Region", "type": "bar", "category_column": "region", "value_column": "revenue", "agg": "sum"},
            {"title": "Revenue Trend", "type": "line", "date_column": "order_date", "value_column": "revenue", "agg": "sum", "freq": "M"},
            {"title": "Bogus chart", "type": "bar", "category_column": "nope", "value_column": "revenue", "agg": "sum"},
        ],
        "insights_text": ["North region drives the largest share of revenue."],
        "risks": ["Revenue is concentrated in a few customers."],
        "opportunities": ["Home category has room to grow."],
        "recommendations": [{"action": "Promote Home category", "priority": "Medium", "impact": "Lift revenue"}],
        "assumptions": ["Discount is stored as a fraction."],
        "limitations": ["No marketing spend data."],
    }


def test_compute_kpi_value_sum(df):
    assert compute_kpi_value(df, {"column": "revenue", "agg": "sum"}) == pytest.approx(1413.25)


def test_compute_kpi_value_row_count(df):
    assert compute_kpi_value(df, {"agg": "row_count"}) == 15


def test_format_kpi_value_currency_small():
    assert format_kpi_value(94.223, "currency") == "$94.22"


def test_format_kpi_value_currency_large():
    assert format_kpi_value(1413.25, "currency") == "$1,413"


def test_format_kpi_value_percent_multiplies_by_100():
    assert format_kpi_value(0.033, "percent") == "3.3%"


def test_format_kpi_value_number_integer():
    assert format_kpi_value(15, "number") == "15"


def test_compute_kpi_summary_drops_invalid_columns(df, plan):
    summary = compute_kpi_summary(df, plan)
    names = [k["name"] for k in summary]
    assert "Bogus" not in names
    assert len(summary) == 3
    assert next(k for k in summary if k["name"] == "Total Revenue")["value"] == "$1,413"
    assert next(k for k in summary if k["name"] == "Total Orders")["value"] == "15"


# ---------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------


def test_build_workbook_has_expected_sheets(df, plan):
    wb = build_workbook(df, plan)
    assert wb.sheetnames == [
        "Dashboard",
        "KPI Reference",
        "Data",
        "Calc",
        "Power BI Guide",
        "Assumptions & Limitations",
        "Lists",
    ]


def test_build_workbook_omits_optional_sheets_when_empty(df):
    plan = {
        "dataset_classification": "Sales",
        "kpis": [{"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency"}],
        "charts": [],
        "filter_columns": [],
    }
    wb = build_workbook(df, plan)
    assert "Assumptions & Limitations" not in wb.sheetnames
    assert "Lists" not in wb.sheetnames
    assert "Dashboard" in wb.sheetnames


def test_build_workbook_is_valid_and_reloadable(df, plan):
    wb = build_workbook(df, plan)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    reloaded = openpyxl.load_workbook(buffer)
    assert reloaded.sheetnames == wb.sheetnames


def test_build_workbook_data_sheet_includes_derived_month_column(df, plan):
    wb = build_workbook(df, plan)
    ws = wb["Data"]
    headers = [c.value for c in ws[1]]
    assert "order_date (Month)" in headers
    assert ws.max_row == len(df) + 1


def test_build_workbook_full_calc_on_load(df, plan):
    wb = build_workbook(df, plan)
    assert wb.calculation.fullCalcOnLoad is True


# ---------------------------------------------------------------------------
# Filters
# ---------------------------------------------------------------------------


def test_build_workbook_drops_invalid_filter_columns(df, plan):
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    # region, category are valid categorical columns; does_not_exist is missing,
    # revenue is numeric so not a valid slicer -> only 2 filters should be wired up
    assert ws["B4"].value == "region:"
    assert ws["E4"].value == "category:"
    assert ws["H4"].value is None


def test_build_workbook_filter_dropdown_has_data_validation(df, plan):
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    assert len(ws.data_validations.dataValidation) == 2
    lists_ws = wb["Lists"]
    assert lists_ws.sheet_state == "hidden"
    assert lists_ws["A1"].value == "All"


def test_filter_dropdown_uses_defined_name_not_raw_cross_sheet_range(df, plan):
    """Data validation lists referencing another sheet directly (formula1="Lists!$A$1:$A$8")
    are a well-documented Excel compatibility trap - some versions/viewers silently
    reject the cross-sheet reference and the dropdown just doesn't work. A workbook
    defined name is the universally-supported way to back a cross-sheet dropdown."""
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    for dv in ws.data_validations.dataValidation:
        assert "!" not in dv.formula1  # bare defined name, not "Sheet!$A$1:$A$8"
        assert dv.formula1 in wb.defined_names
        assert wb.defined_names[dv.formula1].attr_text.startswith("Lists!$")


def test_build_workbook_no_filters_uses_plain_aggregate_formula(df):
    plan = {
        "dataset_classification": "Sales",
        "kpis": [{"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency"}],
        "charts": [],
        "filter_columns": [],
    }
    wb = build_workbook(df, plan)
    calc = wb["Calc"]
    formula = next(c.value for row in calc.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=SUM("))
    assert formula == "=SUM(Data!$J$2:$J$16)"


# ---------------------------------------------------------------------------
# Calc formula engine
# ---------------------------------------------------------------------------


def _all_formula_cells(ws):
    return [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]


def test_calc_sheet_kpi_formulas_reference_data_sheet(df, plan):
    wb = build_workbook(df, plan)
    calc = wb["Calc"]
    formulas = [c.value for c in _all_formula_cells(calc)]
    assert any("SUMIFS(Data!$J$2:$J$16" in f for f in formulas)  # Total Revenue -> revenue column
    assert any(f.startswith("=COUNTIFS(") for f in formulas)  # Total Orders (row_count with filters)


def test_calc_sheet_median_kpi_falls_back_to_static_value(df):
    plan = {
        "dataset_classification": "Sales",
        "filter_columns": ["region"],
        "kpis": [{"name": "Median Revenue", "column": "revenue", "agg": "median", "format": "currency"}],
        "charts": [],
    }
    wb = build_workbook(df, plan)
    calc = wb["Calc"]
    value_cell = next(c for row in calc.iter_rows() for c in row if c.value == df["revenue"].median())
    assert not str(value_cell.value).startswith("=")


def test_dashboard_kpi_cards_reference_calc_sheet(df, plan):
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=Calc!")]
    assert len(formulas) == 3  # 3 valid KPIs (Bogus dropped)


def test_breakdown_table_excludes_own_axis_from_cross_filters(df, plan):
    wb = build_workbook(df, plan)
    calc = wb["Calc"]
    # "Revenue by Region" breaks down by region, so its row formulas must not filter
    # on region again (that would always yield 0), only on the other filter (category).
    region_row_formula = next(
        c.value for row in calc.iter_rows() for c in row
        if isinstance(c.value, str) and "SUMIFS(Data!$J$2:$J$16,Data!$C$2:$C$16,$A" in c.value
    )
    assert "Data!$D$2:$D$16,Calc!" in region_row_formula  # cross-filtered by category
    assert region_row_formula.count("Data!$C$2:$C$16") == 1  # region column only used once (own axis)


def test_diverging_chart_gets_pos_neg_helper_columns(df):
    import pandas as pd

    frame = pd.DataFrame(
        {
            "segment": ["A", "A", "B", "B"],
            "profit": [-500, 200, 1000, 800],
        }
    )
    plan = {
        "dataset_classification": "Finance",
        "filter_columns": [],
        "kpis": [{"name": "Total Profit", "column": "profit", "agg": "sum", "format": "currency"}],
        "charts": [{"title": "Profit by Segment", "type": "bar", "category_column": "segment", "value_column": "profit", "agg": "sum"}],
    }
    wb = build_workbook(frame, plan)
    calc = wb["Calc"]
    # Find the breakdown table header row (has "Pos"/"Neg" columns)
    pos_neg_header = None
    for row in calc.iter_rows():
        values = [c.value for c in row]
        if "Pos" in values and "Neg" in values:
            pos_neg_header = row
            break
    assert pos_neg_header is not None

    ws = wb["Dashboard"]
    chart = ws._charts[0]
    assert type(chart).__name__ == "BarChart"
    assert chart.grouping == "stacked"


def test_doughnut_chart_used_for_low_cardinality_positive_series(df):
    import pandas as pd

    frame = pd.DataFrame({"region": ["North", "South", "East", "West"] * 5, "revenue": [100, 200, 150, 50] * 5})
    plan = {
        "dataset_classification": "Sales",
        "filter_columns": [],
        "kpis": [{"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency"}],
        "charts": [{"title": "Revenue Share", "type": "doughnut", "category_column": "region", "value_column": "revenue", "agg": "sum"}],
    }
    wb = build_workbook(frame, plan)
    ws = wb["Dashboard"]
    assert type(ws._charts[0]).__name__ == "DoughnutChart"


def _spans_overlap(a, b) -> bool:
    (a_col0, a_row0, a_col1, a_row1) = a
    (b_col0, b_row0, b_col1, b_row1) = b
    cols_overlap = a_col0 < b_col1 and b_col0 < a_col1
    rows_overlap = a_row0 < b_row1 and b_row0 < a_row1
    return cols_overlap and rows_overlap


def test_dashboard_charts_use_grid_anchors_with_no_overlap(df, plan):
    """Charts are placed with a TwoCellAnchor bound to exact grid cells (not a
    single-cell anchor + a floating cm width/height, which can visually drift into a
    neighboring chart depending on the sheet's actual column widths)."""
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    assert len(ws._charts) == 2  # `plan` fixture has 2 valid charts (the 3rd is bogus and dropped)

    spans = []
    for chart in ws._charts:
        anchor = chart.anchor
        assert type(anchor).__name__ == "TwoCellAnchor"
        spans.append((anchor._from.col, anchor._from.row, anchor.to.col, anchor.to.row))

    for i in range(len(spans)):
        for j in range(i + 1, len(spans)):
            assert not _spans_overlap(spans[i], spans[j]), (spans[i], spans[j])


def test_dashboard_five_charts_all_slots_no_overlap(df):
    """Exercises every placement slot (trend top-left, doughnut top-right, 3 bottom
    bars) at once - the scenario most likely to overlap if the grid math is wrong."""
    plan = {
        "dataset_classification": "Sales",
        "filter_columns": [],
        "kpis": [{"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency"}],
        "charts": [
            {"title": "Revenue Trend", "type": "line", "date_column": "order_date", "value_column": "revenue", "agg": "sum"},
            {"title": "Revenue Share", "type": "doughnut", "category_column": "region", "value_column": "revenue", "agg": "sum"},
            {"title": "Revenue by Category", "type": "bar", "category_column": "category", "value_column": "revenue", "agg": "sum"},
            {"title": "Revenue by Product", "type": "bar", "category_column": "product", "value_column": "revenue", "agg": "sum"},
            {"title": "Cost by Product", "type": "bar", "category_column": "product", "value_column": "cost", "agg": "sum"},
        ],
    }
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    assert len(ws._charts) == 5

    spans = [(c.anchor._from.col, c.anchor._from.row, c.anchor.to.col, c.anchor.to.row) for c in ws._charts]
    for i in range(len(spans)):
        for j in range(i + 1, len(spans)):
            assert not _spans_overlap(spans[i], spans[j]), (spans[i], spans[j])


def test_chart_series_have_populated_value_cache(df):
    """openpyxl's add_data()/set_categories() only ever write a formula reference to
    the (as-yet uncalculated) Calc cells - no numCache/strCache. A chart with an empty
    cache can render as completely blank in many viewers regardless of whether the
    underlying formula is correct, since they draw from the cache before any
    recalculation happens. The cache must hold the real, pandas-computed numbers."""
    plan = {
        "dataset_classification": "Sales",
        "filter_columns": [],
        "kpis": [{"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency"}],
        "charts": [{"title": "Revenue by Region", "type": "bar", "category_column": "region", "value_column": "revenue", "agg": "sum"}],
    }
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    series_obj = ws._charts[0].series[0]

    assert series_obj.val.numRef.numCache is not None
    cached_values = [pt.v for pt in series_obj.val.numRef.numCache.pt]
    expected = df.groupby("region")["revenue"].sum().sort_values(key=abs, ascending=False)
    assert cached_values == pytest.approx(list(expected.values))


def test_chart_categories_use_str_ref_not_num_ref(df):
    """Category labels here are always text (region/product/month names). openpyxl's
    set_categories() defaults to numRef even for text cells, which is semantically
    wrong and a plausible cause of a blank or broken category axis."""
    plan = {
        "dataset_classification": "Sales",
        "filter_columns": [],
        "kpis": [{"name": "Total Revenue", "column": "revenue", "agg": "sum", "format": "currency"}],
        "charts": [{"title": "Revenue by Region", "type": "bar", "category_column": "region", "value_column": "revenue", "agg": "sum"}],
    }
    wb = build_workbook(df, plan)
    ws = wb["Dashboard"]
    cat = ws._charts[0].series[0].cat
    assert cat.numRef is None
    assert cat.strRef is not None
    cached_labels = [pt.v for pt in cat.strRef.strCache.pt]
    expected_labels = list(df.groupby("region")["revenue"].sum().sort_values(key=abs, ascending=False).index)
    assert cached_labels == expected_labels


def test_diverging_chart_pos_neg_series_have_correct_cached_values():
    frame = pd.DataFrame(
        {
            "segment": ["Enterprise", "Enterprise", "Enterprise", "Government", "Government", "Government", "SMB", "SMB"],
            "profit": [-500, -300, 200, 1000, 800, 1200, 50, -20],
        }
    )
    plan = {
        "dataset_classification": "Finance",
        "filter_columns": [],
        "kpis": [{"name": "Total Profit", "column": "profit", "agg": "sum", "format": "currency"}],
        "charts": [{"title": "Profit by Segment", "type": "bar", "category_column": "segment", "value_column": "profit", "agg": "sum"}],
    }
    wb = build_workbook(frame, plan)
    ws = wb["Dashboard"]
    pos_series, neg_series = ws._charts[0].series

    labels = [pt.v for pt in pos_series.cat.strRef.strCache.pt]
    pos_values = dict(zip(labels, [pt.v for pt in pos_series.val.numRef.numCache.pt]))
    neg_values = dict(zip(labels, [pt.v for pt in neg_series.val.numRef.numCache.pt]))

    assert pos_values["Government"] == pytest.approx(3000)
    assert neg_values["Government"] == pytest.approx(0)
    assert pos_values["Enterprise"] == pytest.approx(0)
    assert neg_values["Enterprise"] == pytest.approx(-600)
    assert pos_values["SMB"] == pytest.approx(30)
    assert neg_values["SMB"] == pytest.approx(0)


# ---------------------------------------------------------------------------
# Formula correctness (independent re-evaluation against pandas ground truth)
# ---------------------------------------------------------------------------


def _resolve_formula(wb, calc, cell) -> float | int:
    """Parse and evaluate one SUMIFS/AVERAGEIFS/COUNTIFS/SUM/COUNTA formula cell against
    the workbook's own Data sheet - an independent formula interpreter, separate from
    dashboard.py's formula-generation code, used to catch bugs a plausible-looking
    formula string could hide (wrong column letter, off-by-one row range, etc.)."""
    import re

    import pandas as pd

    data_ws = wb["Data"]
    headers = [c.value for c in data_ws[1]]
    col_letter_to_name = {openpyxl.utils.get_column_letter(i + 1): h for i, h in enumerate(headers)}
    frame = pd.DataFrame(list(data_ws.iter_rows(min_row=2, values_only=True)), columns=headers)
    range_re = re.compile(r"Data!\$([A-Z]+)\$2:\$([A-Z]+)\$(\d+)")

    def resolve_calc_cell(addr):
        m = re.match(r"Calc!\$([A-Z]+)\$(\d+)", addr)
        v = calc[f"{m.group(1)}{m.group(2)}"].value
        if isinstance(v, str) and v.startswith("="):
            m2 = re.match(r'=IF\(Dashboard!\$([A-Z]+)\$(\d+)="All","\*",Dashboard!\$([A-Z]+)\$(\d+)\)', v)
            dash_val = wb["Dashboard"][f"{m2.group(1)}{m2.group(2)}"].value
            return "*" if dash_val == "All" else dash_val
        return v

    v = cell.value
    if v.startswith("=SUM("):
        col = col_letter_to_name[range_re.search(v).group(1)]
        return float(pd.to_numeric(frame[col], errors="coerce").sum())
    if v.startswith("=COUNTA("):
        col = col_letter_to_name[range_re.search(v).group(1)]
        return int(frame[col].notna().sum())

    func_match = re.match(r"=(\w+)\((.*)\)$", v)
    func, args = func_match.group(1), [a.strip() for a in func_match.group(2).split(",")]
    own_row_values = {}
    label = calc[f"A{cell.row}"].value
    if label is not None:
        own_row_values["A"] = label

    def resolve_criterion(token):
        m = re.match(r"^\$([A-Z]+)(\d+)$", token)
        if m:
            return own_row_values[m.group(1)]
        if token.startswith("Calc!"):
            return resolve_calc_cell(token)
        return token

    start = 0 if func == "COUNTIFS" else 1
    mask = pd.Series(True, index=frame.index)
    for i in range(start, len(args), 2):
        col_name = col_letter_to_name[range_re.search(args[i]).group(1)]
        crit = resolve_criterion(args[i + 1])
        if crit != "*":
            mask &= frame[col_name].astype(str) == str(crit)

    if func == "COUNTIFS":
        return int(mask.sum())
    value_col = col_letter_to_name[range_re.search(args[0]).group(1)]
    values = pd.to_numeric(frame[value_col], errors="coerce")
    subset = values[mask]
    return float(subset.sum() if func == "SUMIFS" else subset.mean())


def test_generated_formulas_evaluate_to_correct_pandas_values(df, plan):
    """Independently interpret every generated formula cell and compare it against a
    directly-computed pandas ground truth for known KPIs/rows - catches bugs a
    plausible-looking formula string could hide (wrong column, wrong row range)."""
    wb = build_workbook(df, plan)
    calc = wb["Calc"]

    expectations = {
        "Total Revenue": df["revenue"].sum(),
        "Total Orders": len(df),
        "Avg Discount": df["discount"].mean(),
    }

    checked = set()
    for row in calc.iter_rows():
        label_cell, value_cell = row[0], row[1]
        if label_cell.value in expectations and isinstance(value_cell.value, str) and value_cell.value.startswith("="):
            result = _resolve_formula(wb, calc, value_cell)
            assert result == pytest.approx(expectations[label_cell.value]), label_cell.value
            checked.add(label_cell.value)

    assert checked == set(expectations)

    # One breakdown-table row, independently: "North" region's revenue with the
    # category cross-filter left at its default "All" (wildcard) state.
    north_row_formula = next(
        c.value
        for r in calc.iter_rows()
        for c in r
        if c.column == 2 and calc.cell(row=c.row, column=1).value == "North"
    )
    assert north_row_formula.startswith("=SUMIFS(")
    north_cell = next(c for r in calc.iter_rows() for c in r if c.value == north_row_formula)
    result = _resolve_formula(wb, calc, north_cell)
    assert result == pytest.approx(df.loc[df["region"] == "North", "revenue"].sum())
